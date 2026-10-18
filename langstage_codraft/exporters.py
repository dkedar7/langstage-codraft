"""Office and motion exports: DOCX, PPTX, GIF and MP4.

These build on the same render as PDF/PNG. The renderer reports where every block
landed on every page (``Renderer.structure``), with element screenshots for anything
that isn't plain text, so:

* **PPTX**: one slide per page, editable text boxes at the exact positions the
  paginator chose; charts, images, tables and split text fragments become pictures.
* **DOCX**: a flowing Word document (Word re-paginates): native headings,
  paragraphs, lists, tables, page size, margins and columns; charts and images are
  pictures.
* **GIF / MP4**: a page-by-page slideshow of the rendered pages. MP4 needs the
  ``[video]`` extra (imageio + a bundled ffmpeg).
"""
from __future__ import annotations

import io
import re
from typing import Any

from .fonts import family_of
from .pages import page_mm
from .store import Store
from .themes import effective_theme


def qn(tag: str) -> str:
    from docx.oxml.ns import qn as _qn

    return _qn(tag)

EMU_PER_PX = 9525  # 914400 EMU/inch at 96 px/inch


# ------------------------------------------------------------------ inline markdown
_INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*]+\*|_[^_]+_|`[^`]+`|\[[^\]]+\]\([^)]+\)|~~[^~]+~~)")


def md_runs(text: str) -> list[tuple[str, dict]]:
    """Split one line of Markdown into (text, style) runs: bold, italic, code, link, strike."""
    runs: list[tuple[str, dict]] = []
    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith(("**", "__")) and len(part) > 4:
            runs.append((part[2:-2], {"bold": True}))
        elif part.startswith("~~"):
            runs.append((part[2:-2], {"strike": True}))
        elif part.startswith("`"):
            runs.append((part[1:-1], {"code": True}))
        elif part.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            runs.append((m.group(1), {"link": m.group(2)}) if m else (part, {}))
        elif part[0] in "*_" and part[-1] == part[0] and len(part) > 2:
            runs.append((part[1:-1], {"italic": True}))
        else:
            runs.append((part, {}))
    return runs


def md_blocks(md: str) -> list[tuple[str, str]]:
    """Coarse Markdown block parse -> [(kind, text)], kind in p, h1-h3, ul, ol, quote, code, hr."""
    out: list[tuple[str, str]] = []
    para: list[str] = []
    code: list[str] | None = None

    def flush():
        if para:
            out.append(("p", " ".join(s.strip() for s in para)))
            para.clear()

    for line in (md or "").splitlines():
        if code is not None:
            if line.strip().startswith("```"):
                out.append(("code", "\n".join(code)))
                code = None
            else:
                code.append(line)
            continue
        s = line.strip()
        if s.startswith("```"):
            flush()
            code = []
        elif not s:
            flush()
        elif re.match(r"^#{1,6} ", s):
            flush()
            level = min(3, len(s) - len(s.lstrip("#")))
            out.append((f"h{level}", s.lstrip("#").strip()))
        elif re.match(r"^[-*+] ", s):
            flush()
            out.append(("ul", s[2:].strip()))
        elif re.match(r"^\d+[.)] ", s):
            flush()
            out.append(("ol", re.sub(r"^\d+[.)] ", "", s)))
        elif s.startswith(">"):
            flush()
            out.append(("quote", s.lstrip("> ").strip()))
        elif re.fullmatch(r"[-*_]{3,}", s):
            flush()
            out.append(("hr", ""))
        else:
            para.append(s)
    flush()
    if code is not None:
        out.append(("code", "\n".join(code)))
    return out


def plain(md: str) -> str:
    return " ".join(t for t, _ in md_runs(md))


# -------------------------------------------------------------------------- DOCX
def to_docx(store: Store, doc_id: str, pictures: dict[str, bytes], references: list[dict]) -> bytes:
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_BREAK
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    doc = store.get_document(doc_id)
    s = doc["settings"]
    blocks = store.list_blocks(doc_id)
    d = docx.Document()
    sec = d.sections[0]
    w, h = page_mm(s)
    sec.page_width, sec.page_height = Mm(w), Mm(h)
    sec.orientation = WD_ORIENT.LANDSCAPE if s.get("orientation") == "landscape" else WD_ORIENT.PORTRAIT
    m = s.get("margins", {})
    sec.top_margin, sec.right_margin = Mm(m.get("top", 20)), Mm(m.get("right", 20))
    sec.bottom_margin, sec.left_margin = Mm(m.get("bottom", 20)), Mm(m.get("left", 20))
    d.core_properties.title = doc["title"]
    content_w_mm = w - m.get("left", 20) - m.get("right", 20)
    theme = effective_theme(s)
    _docx_styles(d, theme)

    def set_columns(section, n: int):
        cols = section._sectPr.find(qn("w:cols"))
        if cols is None:
            cols = section._sectPr.makeelement(qn("w:cols"), {})
            section._sectPr.append(cols)
        cols.set(qn("w:num"), str(n))
        cols.set(qn("w:space"), str(int(s.get("column_gap", 8) / 25.4 * 1440)))

    def add_runs(par, text: str, base: dict | None = None):
        for t, st in md_runs(text):
            r = par.add_run(t)
            st = {**(base or {}), **st}
            r.bold = st.get("bold") or None
            r.italic = st.get("italic") or None
            if st.get("strike"):
                r.font.strike = True
            if st.get("code"):
                r.font.name = "Consolas"
            if st.get("link"):
                r.font.color.rgb = _rgb(theme.get("color_accent")) or RGBColor(0x1A, 0x4D, 0x8F)
                r.underline = True

    def add_markdown(md: str, style_quote=False):
        for kind, text in md_blocks(md):
            if kind.startswith("h"):
                d.add_heading(text, level=min(3, int(kind[1]) + 1))
            elif kind == "ul":
                add_runs(d.add_paragraph(style="List Bullet"), text)
            elif kind == "ol":
                add_runs(d.add_paragraph(style="List Number"), text)
            elif kind == "quote" or style_quote:
                add_runs(d.add_paragraph(style="Quote"), text)
            elif kind == "code":
                p = d.add_paragraph()
                r = p.add_run(text)
                r.font.name = "Consolas"
                r.font.size = Pt(9)
            elif kind == "hr":
                d.add_paragraph("―" * 20)
            else:
                add_runs(d.add_paragraph(), text)

    def picture(bid: str, caption: str):
        data = pictures.get(bid)
        if data:
            d.add_picture(io.BytesIO(data), width=Mm(min(content_w_mm, 170 if s.get("columns", 1) == 1 else content_w_mm / s["columns"])))
        if caption:
            p = d.add_paragraph(caption)
            p.style = d.styles["Caption"]

    columns = int(s.get("columns", 1)) if s.get("mode") == "paged" else 1
    cover = [b for b in blocks if b["layout"].get("role") == "cover"]
    flow = [b for b in blocks if b["layout"].get("role", "flow") == "flow"]
    back = [b for b in blocks if b["layout"].get("role") == "back"]

    def emit(b: dict):
        c, t = b["content"], b["type"]
        if t == "heading":
            d.add_heading(c.get("text", ""), level=max(0, min(3, int(c.get("level", 1)))) if b["layout"].get("role") != "cover" else 0)
        elif t == "text":
            add_markdown(c.get("markdown", ""))
        elif t == "callout":
            add_markdown(c.get("markdown", ""), style_quote=True)
        elif t == "table":
            rows = c.get("rows") or []
            if rows:
                n = max(len(r) for r in rows)
                tbl = d.add_table(rows=len(rows), cols=n)
                tbl.style = "Table Grid"
                for i, r in enumerate(rows):
                    for j in range(n):
                        cell = tbl.cell(i, j)
                        cell.text = ""
                        add_runs(cell.paragraphs[0], str(r[j]) if j < len(r) else "",
                                 {"bold": i == 0 and c.get("header", True)})
            if c.get("caption"):
                d.add_paragraph(c["caption"]).style = d.styles["Caption"]
        elif t in ("image", "chart", "video"):
            picture(b["id"], c.get("caption", ""))
        elif t == "divider":
            d.add_paragraph("―" * 20)
        elif t == "references":
            d.add_heading(c.get("title", "References"), level=2)
            for r in references:
                p = d.add_paragraph(style="List Number")
                add_runs(p, r["title"] + (f". {r['url']}" if r.get("url") else ""))
        elif t == "page-break":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    separate_cover = bool(cover) and (s.get("cover") or {}).get("enabled", True)
    for b in cover:
        emit(b)
    if separate_cover:
        d.add_section()  # new page; flow section gets the columns
    if columns > 1:
        set_columns(d.sections[-1], columns)
    flow_section = d.sections[-1]
    if separate_cover:
        _clear_header_footer(d.sections[0])
    _header_footer(flow_section, s, doc["title"], content_w_mm)
    for b in flow:
        emit(b)
    if back:
        from docx.enum.section import WD_SECTION

        sec_back = d.add_section(WD_SECTION.NEW_PAGE)
        set_columns(sec_back, 1)
        _clear_header_footer(sec_back, unlink=True)
        for b in back:
            emit(b)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _rgb(color: str | None) -> Any:
    from docx.shared import RGBColor

    m = re.fullmatch(r"#([0-9a-fA-F]{6})", (color or "").strip())
    return RGBColor.from_string(m.group(1).upper()) if m else None


def _docx_styles(d: Any, theme: dict) -> None:
    """The document's fonts, sizes and colours on Word's Normal, Title and Heading styles."""
    from docx.shared import Pt

    body_pt = float(theme.get("font_size", 15)) * 0.75
    scale = float(theme.get("heading_scale", 1.25))
    body, head = family_of(theme.get("font_body", "")), family_of(theme.get("font_heading", ""))
    display = family_of(theme.get("font_display") or "") or head
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = body or None, Pt(round(body_pt, 1))
    if _rgb(theme.get("color_text")):
        normal.font.color.rgb = _rgb(theme.get("color_text"))
    heading_color = theme.get("color_heading")
    for name, size, color, family in (
        ("Title", body_pt * scale ** 4, heading_color or theme.get("color_text"), display),
        ("Heading 1", body_pt * scale ** 3, heading_color or theme.get("color_text"), head),
        ("Heading 2", body_pt * scale ** 2, heading_color or theme.get("color_accent"), head),
        ("Heading 3", body_pt * scale, heading_color or theme.get("color_text"), head),
    ):
        try:
            st = d.styles[name]
        except KeyError:
            continue
        st.font.name, st.font.size = family or None, Pt(round(size, 1))
        st.font.bold = int(theme.get("heading_weight", 700)) >= 600
        if _rgb(color):
            st.font.color.rgb = _rgb(color)
        rpr = st.element.get_or_add_rPr()
        fonts_el = rpr.find(qn("w:rFonts"))
        if fonts_el is not None:  # drop theme-font links so the explicit family wins
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                fonts_el.attrib.pop(qn(attr), None)


def _field(run: Any, instr: str) -> None:
    """A Word field (PAGE, NUMPAGES) inside a run."""
    from docx.oxml import OxmlElement

    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        else:
            el = OxmlElement("w:t")
            el.text = text
        run._r.append(el)


def _clear_header_footer(section: Any, unlink: bool = False) -> None:
    for part in (section.header, section.footer):
        if unlink:
            part.is_linked_to_previous = False
        for par in part.paragraphs:
            for r in list(par.runs):
                r._r.getparent().remove(r._r)


def _header_footer(section: Any, s: dict, title: str, content_w_mm: float) -> None:
    """The running header and footer: left/centre/right on tab stops, {page} and {pages}
    as live Word fields."""
    from datetime import datetime

    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.shared import Mm, Pt

    first_hidden = False
    for part_name in ("header", "footer"):
        hf = s.get(part_name) or {}
        slots = [hf.get(k) or "" for k in ("left", "center", "right")]
        target = getattr(section, part_name)
        target.is_linked_to_previous = False
        if not any(slots):
            continue
        first_hidden = first_hidden or not hf.get("first_page", True)
        par = target.paragraphs[0]
        par.paragraph_format.tab_stops.add_tab_stop(Mm(content_w_mm / 2), WD_TAB_ALIGNMENT.CENTER)
        par.paragraph_format.tab_stops.add_tab_stop(Mm(content_w_mm), WD_TAB_ALIGNMENT.RIGHT)
        for i, text in enumerate(slots):
            if i:
                par.add_run("\t")
            text = text.replace("{title}", title).replace("{year}", str(datetime.now().year))
            for piece in re.split(r"(\{page\}|\{pages\})", text):
                if piece in ("{page}", "{pages}"):
                    _field(par.add_run(), "PAGE" if piece == "{page}" else "NUMPAGES")
                elif piece:
                    par.add_run(piece)
        for r in par.runs:
            r.font.size = Pt(8)
        if hf.get("rule"):
            ppr = par._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            edge = OxmlElement("w:bottom" if part_name == "header" else "w:top")
            for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "4"), ("w:color", "999999")):
                edge.set(qn(k), v)
            bdr.append(edge)
            ppr.append(bdr)
    if first_hidden:
        section.different_first_page_header_footer = True


# -------------------------------------------------------------------------- PPTX
def to_pptx(store: Store, doc_id: str, structure: dict) -> bytes:
    from pptx import Presentation
    from pptx.util import Emu, Pt

    doc = store.get_document(doc_id)
    blocks = {b["id"]: b for b in store.list_blocks(doc_id)}
    prs = Presentation()
    pw, ph = structure["page_px"]
    prs.slide_width, prs.slide_height = Emu(int(pw * EMU_PER_PX)), Emu(int(ph * EMU_PER_PX))
    prs.core_properties.title = doc["title"]
    blank = prs.slide_layouts[6]
    theme = structure.get("theme", {})
    body_pt = float(theme.get("font_size", 15)) * 0.75
    scale = float(theme.get("heading_scale", 1.25))

    for page in structure["pages"]:
        slide = prs.slides.add_slide(blank)
        for it in page["items"]:
            b = blocks.get(it["id"])
            x, y, w, h = (Emu(int(v * EMU_PER_PX)) for v in (it["x"], it["y"], it["w"], it["h"]))
            textual = b is not None and b["type"] in ("text", "heading", "callout") and not it["fragment"]
            if textual:
                tb = slide.shapes.add_textbox(x, y, w, h)
                tf = tb.text_frame
                tf.word_wrap = True
                for attr in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
                    setattr(tf, attr, 0)
                first = True
                if b["type"] == "heading":
                    lvl = int(b["content"].get("level", 1))
                    items = [("h", b["content"].get("text", ""), body_pt * scale ** (4 - lvl))]
                else:
                    items = [(k, t, body_pt * (scale ** (4 - int(k[1])) if k in ("h1", "h2", "h3") else 1.0))
                             for k, t in md_blocks(b["content"].get("markdown", ""))]
                for kind, text, size in items:
                    para = tf.paragraphs[0] if first else tf.add_paragraph()
                    first = False
                    prefix = "• " if kind == "ul" else ""
                    for t, st in md_runs(prefix + text):
                        r = para.add_run()
                        r.text = t
                        r.font.size = Pt(size)
                        r.font.bold = True if (st.get("bold") or kind.startswith("h")) else None
                        r.font.italic = st.get("italic") or None
                    para.space_after = Pt(size * 0.5)
            elif it.get("png"):
                slide.shapes.add_picture(io.BytesIO(it["png"]), x, y, w, h)
        if page.get("notes"):
            slide.notes_slide.notes_text_frame.text = page["notes"]
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------- GIF / MP4
def to_gif(pages: list[bytes], seconds_per_page: float = 2.5) -> bytes:
    from PIL import Image

    frames = [Image.open(io.BytesIO(p)).convert("RGB") for p in pages]
    if not frames:
        raise ValueError("no pages to animate")
    pal = [f.convert("P", palette=Image.Palette.ADAPTIVE, colors=255) for f in frames]
    buf = io.BytesIO()
    pal[0].save(buf, format="GIF", save_all=True, append_images=pal[1:], loop=0,
                duration=int(seconds_per_page * 1000), optimize=True)
    return buf.getvalue()


def to_mp4(pages: list[bytes], seconds_per_page: float = 3.0, fps: int = 10) -> bytes:
    """A page-by-page slideshow video (H.264). Needs the ``[video]`` extra."""
    try:
        import imageio
        import numpy as np
    except ImportError as exc:  # pragma: no cover - optional extra
        raise RuntimeError('MP4 export needs the video extra: pip install "langstage-codraft[video]"') from exc
    import tempfile
    from pathlib import Path

    from PIL import Image

    imgs = [Image.open(io.BytesIO(p)).convert("RGB") for p in pages]
    if not imgs:
        raise ValueError("no pages to animate")
    w = max(i.width for i in imgs) // 2 * 2
    h = max(i.height for i in imgs) // 2 * 2
    with tempfile.TemporaryDirectory() as td:
        out = Path(td) / "doc.mp4"
        writer = imageio.get_writer(str(out), fps=fps, codec="libx264", quality=8, macro_block_size=2)
        try:
            for im in imgs:
                canvas = Image.new("RGB", (w, h), "white")
                canvas.paste(im, ((w - im.width) // 2, (h - im.height) // 2))
                arr = np.asarray(canvas)
                for _ in range(max(1, int(seconds_per_page * fps))):
                    writer.append_data(arr)
        finally:
            writer.close()
        return out.read_bytes()
