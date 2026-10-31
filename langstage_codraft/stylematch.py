"""Read the look of an uploaded document, so a new one can match it.

``analyze`` returns a *style profile*: page size, margins and columns; the body, heading
and display fonts (mapped to open families by :mod:`fonts`); text, heading, accent and
tint colours; whether it has a separate cover and what it looks like; and running
header/footer text. ``to_settings`` turns a profile into Codraft document settings.

* PDF: fonts, sizes and colours come straight from the text runs (pypdf); page setup,
  page colours and the cover from the rendered pages (pypdfium2 + Pillow).
* Word (.docx): page setup and the Normal / Heading / Title styles, resolving theme
  fonts and theme colours.
* Page images (PNG/JPG): colours, margins and columns from the pixels; fonts are left
  to the agent, which sees the image.

Everything here is a best guess from a sample. The agent sees the rendered pages (see
``render_pages``) and refines the cover, block styles and anything the numbers miss.
"""
from __future__ import annotations

import io
import math
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from . import fonts
from .pages import PAGE_SIZES_MM

MM_PER_PT = 25.4 / 72
MAX_PAGES = 6        # pages rendered for the agent and sampled for colours
RENDER_DPI = 110
# The named sizes a document can store (gh #30): anything else matches as custom.
STANDARD_SIZES = PAGE_SIZES_MM


# ------------------------------------------------------------------ rendering
def render_pages(data: bytes, out_dir: Path, stem: str, max_pages: int = MAX_PAGES) -> list[str]:
    """Render the first pages of a PDF to JPEGs (for the agent's eyes and for analysis)."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    paths = []
    try:
        for i in range(min(len(pdf), max_pages)):
            page = pdf[i]
            img = page.render(scale=RENDER_DPI / 72).to_pil().convert("RGB")
            path = out_dir / f"{stem}-p{i + 1}.jpg"
            img.save(path, "JPEG", quality=82)
            paths.append(str(path))
            page.close()
    finally:
        pdf.close()
    return paths


def _pdf_images(data: bytes, max_pages: int = MAX_PAGES, dpi: int = 60) -> list[Any]:
    """Lossless in-memory renders for colour analysis (JPEG artefacts shift flat colours)."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    out = []
    try:
        for i in range(min(len(pdf), max_pages)):
            page = pdf[i]
            out.append(page.render(scale=dpi / 72).to_pil().convert("RGB"))
            page.close()
    finally:
        pdf.close()
    return out


# ------------------------------------------------------------------ colour helpers
def _hex(rgb: tuple[float, float, float]) -> str:
    return "#" + "".join(f"{max(0, min(255, round(c))):02x}" for c in rgb)


def _rgb(h: str) -> tuple[int, int, int]:
    h = h.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def luminance(h: str) -> float:
    r, g, b = (c / 255 for c in _rgb(h))
    lin = [c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4 for c in (r, g, b)]
    return 0.2126 * lin[0] + 0.7152 * lin[1] + 0.0722 * lin[2]


def chroma(h: str) -> float:
    r, g, b = _rgb(h)
    return (max(r, g, b) - min(r, g, b)) / 255


def _dist(a: str, b: str) -> float:
    return math.dist(_rgb(a), _rgb(b))


def _cmyk(c: float, m: float, y: float, k: float) -> tuple[float, float, float]:
    return (255 * (1 - c) * (1 - k), 255 * (1 - m) * (1 - k), 255 * (1 - y) * (1 - k))


# ------------------------------------------------------------------ PDF text runs
def _num(v: Any) -> float | None:
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _pdf_runs(data: bytes, max_pages: int = 12) -> tuple[list[dict], list[tuple[float, float]]]:
    """Text runs (font, size pt, colour, position from the top-left in pt, chars) per page."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    runs: list[dict] = []
    sizes: list[tuple[float, float]] = []
    for pi, page in enumerate(reader.pages[:max_pages]):
        box = page.mediabox
        pw, ph = float(box.width), float(box.height)
        sizes.append((pw, ph))
        state = {"fill": (0.0, 0.0, 0.0), "stack": []}

        def before(op: bytes, args: list, cm: list, tm: list, state: dict = state) -> None:
            nums = [_num(a) for a in args]
            if op == b"q":
                state["stack"].append(state["fill"])
            elif op == b"Q" and state["stack"]:
                state["fill"] = state["stack"].pop()
            elif op == b"rg" and len(nums) == 3 and None not in nums:
                state["fill"] = tuple(255 * n for n in nums)
            elif op == b"g" and len(nums) == 1 and nums[0] is not None:
                state["fill"] = (255 * nums[0],) * 3
            elif op == b"k" and len(nums) == 4 and None not in nums:
                state["fill"] = _cmyk(*nums)
            elif op in (b"sc", b"scn") and nums and None not in nums:
                if len(nums) == 3:
                    state["fill"] = tuple(255 * n for n in nums)
                elif len(nums) == 1:
                    state["fill"] = (255 * nums[0],) * 3
                elif len(nums) == 4:
                    state["fill"] = _cmyk(*nums)

        def text(t: str, cm: list, tm: list, font: Any, size: Any, state: dict = state,
                 pi: int = pi, ph: float = ph) -> None:
            chars = len(t.strip())
            if not chars or not font or size is None:
                return
            scale_tm = math.hypot(tm[2], tm[3]) or 1.0
            scale_cm = math.hypot(cm[2], cm[3]) or 1.0
            pt = float(size) * scale_tm * scale_cm
            if not 3 <= pt <= 200:
                return
            x = cm[0] * tm[4] + cm[2] * tm[5] + cm[4]
            y = cm[1] * tm[4] + cm[3] * tm[5] + cm[5]
            name = str(font.get("/BaseFont", "")) if hasattr(font, "get") else ""
            flags = 0
            try:
                desc = font.get("/FontDescriptor")
                desc = desc.get_object() if desc is not None else None
                flags = int(desc.get("/Flags", 0)) if desc is not None else 0
            except Exception:  # noqa: BLE001 - a malformed descriptor only costs the hint
                flags = 0
            runs.append({"page": pi, "font": name, "size": round(pt * 2) / 2, "chars": chars,
                         "color": _hex(state["fill"]), "x": x, "top": ph - y, "text": t.strip(),
                         "serif": bool(flags & 2), "mono": bool(flags & 1)})

        try:
            page.extract_text(visitor_operand_before=before, visitor_text=text)
        except Exception:  # noqa: BLE001 - one odd page shouldn't sink the profile
            continue
    return runs, sizes


def _weighted_mode(items: list[tuple[Any, int]]) -> Any:
    c: Counter = Counter()
    for k, w in items:
        c[k] += w
    return c.most_common(1)[0][0] if c else None


def _fonts_from_runs(runs: list[dict], every: list[dict] | None = None) -> dict[str, Any]:
    """Fonts, sizes and colours from body-page runs; the display face from all pages."""
    every = every or runs
    body_size = _weighted_mode([(r["size"], r["chars"]) for r in runs])
    if not body_size:
        return {}
    body_runs = [r for r in runs if abs(r["size"] - body_size) <= 0.6]
    body_font = _weighted_mode([(r["font"], r["chars"]) for r in body_runs])
    head_runs = [r for r in runs if r["size"] >= body_size * 1.18 and r["chars"] >= 3]
    out: dict[str, Any] = {"body_size_pt": body_size}

    def mapped(raw: str, rs: list[dict]) -> dict:
        serif = any(r["serif"] for r in rs if r["font"] == raw)
        mono = any(r["mono"] for r in rs if r["font"] == raw)
        return fonts.map_font(raw, serif=serif, mono=mono)

    out["body"] = mapped(body_font, body_runs)
    out["text_color"] = _weighted_mode([(r["color"], r["chars"]) for r in body_runs])
    if head_runs:
        head_font = _weighted_mode([(r["font"], r["chars"]) for r in head_runs])
        out["heading"] = mapped(head_font, head_runs)
        top = max(r["size"] for r in head_runs)
        top_all = max(r["size"] for r in every)
        big = [r for r in every if r["size"] >= max(body_size * 2, top_all * 0.8) and r["chars"] >= 3]
        if big:
            out["display"] = mapped(_weighted_mode([(r["font"], r["chars"]) for r in big]), big)
            out["title_size_pt"] = top_all
        # The most common heading size stands for h2 (h2 = body * scale^2).
        sec = [r for r in head_runs if r["size"] < top * 0.95] or head_runs
        h2 = _weighted_mode([(r["size"], r["chars"]) for r in sec])
        out["heading_scale"] = round(max(1.05, min(1.6, math.sqrt(h2 / body_size))), 3)
        out["heading_color"] = _weighted_mode([(r["color"], r["chars"]) for r in sec])
        colors = Counter()
        for r in runs:
            if r["color"] not in (out["text_color"], out["heading_color"]) and chroma(r["color"]) > 0.25:
                colors[r["color"]] += r["chars"]
        out["accent_candidates"] = [c for c, _ in colors.most_common(3)]
    # Line height: the usual baseline step between body lines in one column.
    steps = []
    by_page: dict[int, list[dict]] = {}
    for r in body_runs:
        by_page.setdefault(r["page"], []).append(r)
    for rs in by_page.values():
        rs.sort(key=lambda r: (round(r["x"]), r["top"]))
        for a, b in zip(rs, rs[1:], strict=False):
            dy = b["top"] - a["top"]
            if abs(a["x"] - b["x"]) < 2 and body_size * 1.0 <= dy <= body_size * 2.6:
                steps.append(round(dy * 2) / 2)
    if steps:
        out["line_height"] = round(Counter(steps).most_common(1)[0][0] / body_size, 2)
    return out


def _furniture(runs: list[dict], sizes: list[tuple[float, float]], margins_pt: dict | None,
               body_pages: list[int]) -> dict[str, dict[str, str]]:
    """Running header/footer text: runs in the top/bottom margin repeated on most pages."""
    if not margins_pt or len(body_pages) < 2:
        return {}
    found: dict[str, Counter] = {"header": Counter(), "footer": Counter()}
    for pi in body_pages:
        if pi >= len(sizes):
            continue
        pw, ph = sizes[pi]
        slots: dict[tuple[str, str], list[str]] = {}
        for r in runs:
            if r["page"] != pi:
                continue
            if r["top"] < margins_pt["top"] - 2:
                part = "header"
            elif r["top"] > ph - margins_pt["bottom"] + 2:
                part = "footer"
            else:
                continue
            slot = "left" if r["x"] < pw / 3 else "right" if r["x"] > pw * 2 / 3 else "center"
            slots.setdefault((part, slot), []).append(r["text"])
        for (part, slot), texts in slots.items():
            t = " ".join(texts).strip()
            t = t.replace(str(datetime.now().year), "{year}")
            t = re.sub(r"(?<![\w{])\d{1,4}(?![\w}])", "{page}", t)
            found[part][(slot, t)] += 1
    out: dict[str, dict[str, str]] = {}
    need = max(2, math.ceil(len(body_pages) / 2))
    for part, c in found.items():
        slots = {slot: t for (slot, t), n in c.items() if n >= need and t}
        if slots:
            out[part] = {s: slots.get(s, "") for s in ("left", "center", "right")}
    return out


# ------------------------------------------------------------------ pixels
def _palette(img: Any, max_side: int = 900, group: int = 6) -> list[tuple[float, str]]:
    """Colours of a page by share of its area, most common first. Near-identical shades
    (anti-aliasing, JPEG noise) are grouped; each group reports its most common exact
    colour, so flat fills come back exactly."""
    small = img.convert("RGB")
    small.thumbnail((max_side, max_side))
    total = small.width * small.height
    groups: dict[tuple[int, ...], Counter] = {}
    for count, rgb in small.getcolors(maxcolors=total) or []:
        groups.setdefault(tuple(v // group for v in rgb), Counter())[rgb] += count
    out = [(sum(c.values()) / total, _hex(c.most_common(1)[0][0])) for c in groups.values()]
    return sorted(out, reverse=True)


def _ink_profile(img: Any, bg: str) -> tuple[list[float], list[float], int, int]:
    """Row and column ink fractions (0..1) of a page image against its background."""
    from PIL import Image, ImageChops

    g = img.convert("RGB")
    g.thumbnail((600, 800))
    diff = ImageChops.difference(g, Image.new("RGB", g.size, _rgb(bg))).convert("L")
    mask = diff.point(lambda v: 255 if v > 48 else 0)
    w, h = mask.size
    rows = [v / 255 for v in mask.resize((1, h), Image.Resampling.BOX).tobytes()]
    cols = [v / 255 for v in mask.resize((w, 1), Image.Resampling.BOX).tobytes()]
    return rows, cols, w, h


def _bands(profile: list[float], min_ink: float, min_gap: int) -> list[tuple[int, int]]:
    bands, start, gap = [], None, 0
    for i, v in enumerate(profile):
        if v > min_ink:
            if start is None:
                start = i
            gap = 0
            end = i
        elif start is not None:
            gap += 1
            if gap >= min_gap:
                bands.append((start, end))
                start, gap = None, 0
    if start is not None:
        bands.append((start, end))
    return bands


def _page_geometry(img: Any, bg: str, w_mm: float, h_mm: float) -> dict | None:
    """Margins (mm, excluding a running header/footer), header/footer presence, columns."""
    rows, cols, w, h = _ink_profile(img, bg)
    bands = _bands(rows, 0.002, max(3, h // 100))
    if not bands:
        return None
    header = footer = False
    if len(bands) > 1 and bands[0][1] < h * 0.12 and (bands[0][1] - bands[0][0]) < h * 0.035 \
            and bands[1][0] - bands[0][1] > h * 0.015:
        header, bands = True, bands[1:]
    if len(bands) > 1 and bands[-1][0] > h * 0.88 and (bands[-1][1] - bands[-1][0]) < h * 0.035 \
            and bands[-1][0] - bands[-2][1] > h * 0.015:
        footer, bands = True, bands[:-1]
    top, bottom = bands[0][0], bands[-1][1]
    # Columns: ink inside the content rows only.
    from PIL import Image, ImageChops

    g = img.convert("RGB")
    g.thumbnail((600, 800))
    crop = g.crop((0, top, w, bottom + 1))
    diff = ImageChops.difference(crop, Image.new("RGB", crop.size, _rgb(bg))).convert("L")
    mask = diff.point(lambda v: 255 if v > 48 else 0)
    cprof = [v / 255 for v in mask.resize((w, 1), Image.Resampling.BOX).tobytes()]
    xs = [i for i, v in enumerate(cprof) if v > 0.002]
    if not xs:
        return None
    left, right = xs[0], xs[-1]
    if left < w * 0.015 or right > w * 0.985 or top < h * 0.01 or bottom > h * 0.99:
        return None  # bleeds off the page: no margin to read
    gutters = []
    run = None
    for i in range(left, right + 1):
        if cprof[i] <= 0.004:
            run = (run[0], i) if run else (i, i)
        else:
            if run and run[1] - run[0] + 1 >= w * 0.015:
                gutters.append(run)
            run = None
    content_w = right - left
    parts, prev = [], left
    for a, b in gutters:
        parts.append(a - prev)
        prev = b
    parts.append(right - prev)
    columns = len(parts) if 1 < len(parts) <= 3 and min(parts) > content_w * 0.2 else 1
    mm_x, mm_y = w_mm / w, h_mm / h
    out = {"margins": {"top": round(top * mm_y, 1), "bottom": round((h - bottom) * mm_y, 1),
                       "left": round(left * mm_x, 1), "right": round((w - right) * mm_x, 1)},
           "columns": columns, "header": header, "footer": footer}
    if columns > 1:
        out["column_gap"] = round(sum(b - a + 1 for a, b in gutters) / len(gutters) * mm_x, 1)
    return out


def _colors_from_images(images: list[Any], accent_min: float = 0.0005) -> dict[str, Any]:
    """Page colour, cover, text, tint and accent colours from rendered pages."""
    if not images:
        return {}
    pals = [_palette(im) for im in images]
    body_idx = list(range(1, len(images))) if len(images) > 1 else [0]
    page_bg = _weighted_mode([(c, share) for i in body_idx for share, c in pals[i][:1]]) or "#ffffff"
    out: dict[str, Any] = {"page": page_bg}
    first_share, first_bg = pals[0][0]
    cover = None
    if len(images) > 1 and first_share >= 0.35 and _dist(first_bg, page_bg) > 60:
        cover = {"background": first_bg, "text": "light" if luminance(first_bg) < 0.35 else "dark"}
    out["cover"] = cover
    texts, tints, accents = Counter(), Counter(), Counter()
    for i in body_idx:
        # Flat fills (tints) read exactly; text and thin rules are anti-aliased, so they
        # are read on coarse groups and come back approximate.
        for share, c in pals[i]:
            if _dist(c, page_bg) >= 14 and luminance(c) > 0.72 and chroma(c) > 0.04 and share >= 0.008:
                tints[c] += share
        for share, c in _palette(images[i], group=32):
            if _dist(c, page_bg) < 40:
                continue
            if luminance(c) < 0.15 and chroma(c) < 0.25 and share >= 0.002:
                texts[c] += share
            elif chroma(c) > 0.3 and 0.02 < luminance(c) < 0.5 and share >= accent_min:
                accents[c] += share
    out["text"] = texts.most_common(1)[0][0] if texts else None
    out["tint"] = tints.most_common(1)[0][0] if tints else None
    out["accent"] = accents.most_common(1)[0][0] if accents else None
    return out


# ------------------------------------------------------------------ analyzers
def _size_name(w_mm: float, h_mm: float) -> tuple[str, str]:
    orient = "landscape" if w_mm > h_mm else "portrait"
    a, b = sorted((w_mm, h_mm))
    for name, (sw, sh) in STANDARD_SIZES.items():
        x, y = sorted((sw, sh))
        if abs(a - x) <= 3 and abs(b - y) <= 3:
            return name, ("landscape" if name == "slide" and orient == "landscape" else orient)
    return "custom", orient


def analyze_pdf(data: bytes) -> dict[str, Any]:
    runs, sizes = _pdf_runs(data)
    if not sizes:
        return {}
    pw, ph = sizes[0]
    w_mm, h_mm = pw * MM_PER_PT, ph * MM_PER_PT
    images = _pdf_images(data)
    prof: dict[str, Any] = {"source": "pdf", "pages": len(sizes)}
    size, orient = _size_name(w_mm, h_mm)
    prof["page"] = {"size": size, "orientation": orient, "w_mm": round(w_mm, 1), "h_mm": round(h_mm, 1)}
    colors = _colors_from_images(images)
    has_cover = bool(colors.get("cover"))
    body_pages = [i for i in range(len(images)) if not (has_cover and i == 0)] or [0]
    geos = [g for i in body_pages if (g := _page_geometry(images[i], colors.get("page", "#ffffff"), w_mm, h_mm))]
    if geos:
        # Top and left: the usual value. Bottom and right: the fullest page (short last pages
        # and ragged lines would otherwise inflate them).
        prof["margins"] = {k: round(sorted(g["margins"][k] for g in geos)[len(geos) // 2], 1)
                           for k in ("top", "left")}
        prof["margins"].update({k: round(min(g["margins"][k] for g in geos), 1) for k in ("right", "bottom")})
        m = prof["margins"]
        if m["bottom"] > 1.6 * max(m["top"], 15):  # no sampled page ran to the bottom: assume symmetry
            m["bottom"] = m["top"]
        prof["columns"] = Counter(g["columns"] for g in geos).most_common(1)[0][0]
        gaps = [g["column_gap"] for g in geos if g.get("column_gap")]
        if gaps and prof["columns"] > 1:
            prof["column_gap"] = sorted(gaps)[len(gaps) // 2]
    f = _fonts_from_runs([r for r in runs if not (has_cover and r["page"] == 0)] or runs, runs)
    prof["fonts"] = {k: f[k] for k in ("body", "heading", "display") if k in f}
    for k in ("body_size_pt", "heading_scale", "line_height", "title_size_pt"):
        if k in f:
            prof[k] = f[k]
    # Accent: a second brand colour when there is one (headings already have their own).
    accent = next((c for c in [*f.get("accent_candidates", []), colors.get("accent"), f.get("heading_color")]
                   if c and chroma(c) > 0.25), None)
    prof["colors"] = {"text": f.get("text_color"), "heading": f.get("heading_color"), "accent": accent,
                      "page": colors.get("page"), "tint": colors.get("tint")}
    prof["cover"] = colors.get("cover") if has_cover else ({"none": True} if len(sizes) > 1 else None)
    if prof.get("margins"):
        mpt = {k: v / MM_PER_PT for k, v in prof["margins"].items()}
        furn = _furniture(runs, sizes, mpt, [i for i in body_pages if i < len(sizes)])
        prof.update(furn)
    return prof


def _theme_xml(d: Any) -> tuple[dict[str, str], dict[str, str]]:
    """Theme fonts (major/minor) and colours (dk1, accent1, ...) of a .docx."""
    from lxml import etree

    fonts_, colors = {}, {}
    for rel in d.part.rels.values():
        if rel.reltype.endswith("/theme"):
            root = etree.fromstring(rel.target_part.blob)
            ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            for kind in ("major", "minor"):
                el = root.find(f".//a:{kind}Font/a:latin", ns)
                if el is not None and el.get("typeface"):
                    fonts_[kind] = el.get("typeface")
            scheme = root.find(".//a:clrScheme", ns)
            if scheme is not None:
                for child in scheme:
                    tag = etree.QName(child).localname
                    val = None
                    for c in child:
                        val = c.get("val") if c.get("val") and etree.QName(c).localname == "srgbClr" else c.get("lastClr")
                    if val:
                        colors[tag] = "#" + val.lower()
    return fonts_, colors


def analyze_docx(data: bytes) -> dict[str, Any]:
    import docx
    from docx.oxml.ns import qn

    d = docx.Document(io.BytesIO(data))
    sec = d.sections[0]
    prof: dict[str, Any] = {"source": "docx"}
    if sec.page_width and sec.page_height:
        w_mm, h_mm = sec.page_width.mm, sec.page_height.mm
        size, orient = _size_name(w_mm, h_mm)
        prof["page"] = {"size": size, "orientation": orient, "w_mm": round(w_mm, 1), "h_mm": round(h_mm, 1)}
    prof["margins"] = {k: round(getattr(sec, f"{k}_margin").mm, 1) for k in ("top", "right", "bottom", "left")
                       if getattr(sec, f"{k}_margin") is not None}
    cols = sec._sectPr.find(qn("w:cols"))
    if cols is not None and cols.get(qn("w:num")):
        prof["columns"] = max(1, min(3, int(cols.get(qn("w:num")))))
        if cols.get(qn("w:space")):
            prof["column_gap"] = round(int(cols.get(qn("w:space"))) / 1440 * 25.4, 1)
    theme_fonts, theme_colors = _theme_xml(d)

    def style_info(name: str) -> dict[str, Any]:
        try:
            st = d.styles[name]
        except KeyError:
            return {}
        info: dict[str, Any] = {}
        cur = st
        while cur is not None and not ("font" in info and "size" in info and "color" in info):
            f = cur.font
            if "font" not in info:
                rpr = cur.element.rPr
                rf = rpr.rFonts if rpr is not None else None
                if f.name:
                    info["font"] = f.name
                elif rf is not None and rf.get(qn("w:asciiTheme")):
                    kind = "major" if "major" in rf.get(qn("w:asciiTheme")) else "minor"
                    if theme_fonts.get(kind):
                        info["font"] = theme_fonts[kind]
            if "size" not in info and f.size:
                info["size"] = f.size.pt
            if "color" not in info and f.color is not None and f.color.type is not None:
                if f.color.rgb is not None:
                    info["color"] = "#" + str(f.color.rgb).lower()
                elif f.color.theme_color is not None:
                    key = {"ACCENT_1": "accent1", "ACCENT_2": "accent2", "TEXT_1": "dk1", "TEXT_2": "dk2",
                           "DARK_1": "dk1", "DARK_2": "dk2"}.get(f.color.theme_color.name)
                    if key and theme_colors.get(key):
                        info["color"] = theme_colors[key]
            if "bold" not in info and f.bold is not None:
                info["bold"] = bool(f.bold)
            cur = cur.base_style
        return info

    normal, h1, h2, title = (style_info(n) for n in ("Normal", "Heading 1", "Heading 2", "Title"))
    body_font = normal.get("font") or theme_fonts.get("minor") or "Calibri"
    head_font = h1.get("font") or h2.get("font") or theme_fonts.get("major") or body_font
    prof["fonts"] = {"body": fonts.map_font(body_font), "heading": fonts.map_font(head_font)}
    if title.get("font") and fonts.map_font(title["font"])["family"] not in (
            prof["fonts"]["body"]["family"], prof["fonts"]["heading"]["family"]):
        prof["fonts"]["display"] = fonts.map_font(title["font"])
    body_pt = normal.get("size") or 11.0
    prof["body_size_pt"] = body_pt
    if h2.get("size"):
        prof["heading_scale"] = round(max(1.05, min(1.6, math.sqrt(h2["size"] / body_pt))), 3)
    text = normal.get("color") or theme_colors.get("dk1")
    heading = h1.get("color") or h2.get("color")
    accent = theme_colors.get("accent1")
    prof["colors"] = {"text": text if text and text != "#000000" else None, "heading": heading,
                      "accent": accent, "page": None, "tint": None}
    return prof


def analyze_image(path: str) -> dict[str, Any]:
    from PIL import Image

    img = Image.open(path)
    ratio = img.width / img.height
    guess = min(STANDARD_SIZES.items(), key=lambda kv: min(abs(kv[1][0] / kv[1][1] - ratio),
                                                           abs(kv[1][1] / kv[1][0] - ratio)))
    w_mm, h_mm = guess[1] if ratio < 1 or guess[0] == "slide" else guess[1][::-1]
    if (w_mm > h_mm) != (ratio > 1):
        w_mm, h_mm = h_mm, w_mm
    prof: dict[str, Any] = {"source": "image"}
    size, orient = _size_name(w_mm, h_mm)
    prof["page"] = {"size": size, "orientation": orient, "w_mm": w_mm, "h_mm": h_mm, "guessed": True}
    colors = _colors_from_images([img], accent_min=0.004)  # sub-pixel fringes look like accents
    geo = _page_geometry(img, colors.get("page", "#ffffff"), w_mm, h_mm)
    if geo:
        prof["margins"] = geo["margins"]
        prof["columns"] = geo["columns"]
        if geo.get("column_gap"):
            prof["column_gap"] = geo["column_gap"]
    prof["colors"] = {"text": colors.get("text"), "heading": None, "accent": colors.get("accent"),
                      "page": colors.get("page"), "tint": colors.get("tint")}
    prof["fonts"] = {}
    return prof


def analyze(filename: str, data: bytes, mime: str, page_paths: list[str]) -> dict[str, Any] | None:
    """Style profile of an upload, or None for a format we can't read."""
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".pdf" or mime == "application/pdf":
            return analyze_pdf(data)
        if ext == ".docx":
            return analyze_docx(data)
        if mime.startswith("image/") and not mime.endswith("svg+xml"):
            return analyze_image(page_paths[0]) if page_paths else None
    except Exception as exc:  # noqa: BLE001 - a file we can't read is a gap, not a crash
        return {"source": ext.lstrip(".") or "file", "error": f"{type(exc).__name__}: {exc}"}
    return None


# ------------------------------------------------------------------ applying
def _px(pt: float) -> float:
    return round(pt * 96 / 72 * 2) / 2


def to_settings(profile: dict[str, Any]) -> tuple[dict[str, Any], list[str]]:
    """Codraft settings that reproduce a profile, and a plain-language list of changes."""
    s: dict[str, Any] = {}
    changes: list[str] = []
    page = profile.get("page") or {}
    if page and not page.get("guessed"):
        if page["size"] == "custom":
            s.update(page_size="custom", page_w_mm=page["w_mm"], page_h_mm=page["h_mm"])
        else:
            s["page_size"] = page["size"]
        s["orientation"] = page["orientation"]
    m = profile.get("margins")
    if m and all(k in m for k in ("top", "right", "bottom", "left")):
        s["margins"] = {k: max(5.0, min(60.0, round(v * 2) / 2)) for k, v in m.items()}
    if profile.get("columns"):
        s["columns"] = int(profile["columns"])
        if profile.get("column_gap") and s["columns"] > 1:
            s["column_gap"] = max(2.0, min(20.0, round(profile["column_gap"] * 2) / 2))
    bits = []
    if "page_size" in s:
        bits.append(f"{'custom ' + str(page['w_mm']) + ' x ' + str(page['h_mm']) + ' mm' if s['page_size'] == 'custom' else s['page_size'].upper() if len(s['page_size']) == 2 else s['page_size'].title()} {s['orientation']}")
    if "margins" in s:
        mm = s["margins"]
        bits.append(f"margins {mm['top']:g}/{mm['right']:g}/{mm['bottom']:g}/{mm['left']:g} mm")
    if "columns" in s:
        bits.append(f"{s['columns']} column{'s' if s['columns'] > 1 else ''}")
    if bits:
        changes.append("Page: " + ", ".join(bits))

    style: dict[str, Any] = {}
    f = profile.get("fonts") or {}
    font_bits = []
    for key, target in (("body", "font_body"), ("heading", "font_heading"), ("display", "font_display")):
        info = f.get(key)
        if not info:
            continue
        style[target] = fonts.css_stack(info["family"], info["category"])
        same = info["exact"] or info["original"].lower() == info["family"].lower()
        font_bits.append(f"{info['family']} ({key})" if same else f"{info['original']} → {info['family']} ({key})")
    if f.get("heading") and f["heading"].get("weight", 400) >= 500:
        style["heading_weight"] = int(f["heading"]["weight"])
    if f.get("heading") and "font_label" not in style:
        style["font_label"] = style["font_heading"]
    if profile.get("body_size_pt"):
        style["font_size"] = max(8.0, min(24.0, _px(profile["body_size_pt"])))
    if profile.get("heading_scale"):
        style["heading_scale"] = profile["heading_scale"]
    if profile.get("line_height"):
        style["line_height"] = max(1.1, min(2.2, profile["line_height"]))
    if font_bits:
        changes.append("Fonts: " + ", ".join(font_bits))
    c = profile.get("colors") or {}
    color_bits = []
    for key, target in (("text", "color_text"), ("heading", "color_heading"), ("accent", "color_accent")):
        if c.get(key):
            style[target] = c[key]
            color_bits.append(f"{key} {c[key]}")
    if c.get("accent"):
        style["color_rule"] = c["accent"]
    if c.get("tint"):
        style["color_callout"] = style["color_figure"] = c["tint"]
        color_bits.append(f"tint {c['tint']}")
    if c.get("page") and luminance(c["page"]) < 0.93:
        style["color_page"] = c["page"]
        color_bits.append(f"paper {c['page']}")
    if color_bits:
        changes.append("Colours: " + ", ".join(color_bits))
    if style:
        s["style"] = style
    cover = profile.get("cover")
    if cover and cover.get("none"):
        s["cover"] = {"enabled": False}
        changes.append("Cover: none (the title opens page 1)")
    elif cover:
        s["cover"] = {"enabled": True, "background": cover["background"], "text": cover["text"]}
        changes.append(f"Cover: {cover['background']} with {cover['text']} text")
    for part in ("header", "footer"):
        if profile.get(part):
            s[part] = {**profile[part], "first_page": True}
            changes.append(f"{part.title()}: " + " | ".join(v or "—" for v in profile[part].values()))
    return s, changes


def summary(profile: dict[str, Any] | None) -> str:
    """One paragraph for the agent."""
    if not profile:
        return "No style profile (format not analysed)."
    if profile.get("error"):
        return f"Could not read its style: {profile['error']}"
    _, changes = to_settings(profile)
    head = {"pdf": "From the PDF's text runs and rendered pages", "docx": "From the Word styles",
            "image": "From the page image (approximate; judge fonts and colours from the image)"}.get(profile.get("source"), "")
    return f"{head}: " + "; ".join(changes) if changes else f"{head}: nothing readable."
