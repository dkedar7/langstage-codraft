"""0.5: cover pages, per-document style, running header/footer, saved themes and matching
an uploaded document's look."""
import io
import json

import pytest
from fastapi.testclient import TestClient

from langstage_codraft import fonts, stylematch, themes
from langstage_codraft import templates as tpl
from langstage_codraft.config import CodraftConfig
from langstage_codraft.server import create_app
from langstage_codraft.store import Invalid, Store
from langstage_codraft.tools import (
    get_document,
    match_upload_style,
    propose_insert,
    set_cover,
    set_header_footer,
    set_typography,
)


def _report(store):
    return store.create_document("Pilot", {}, tpl.BUILTIN_TEMPLATES["report"]["blocks"])


def _app(tmp_path, **kw):
    cfg = CodraftConfig.resolve(overrides={"data_dir": tmp_path / "data", "port": 8797}, use_toml=False)
    return create_app(cfg, None, {"kind": "none"}, store=Store(tmp_path / "data" / "codraft.db"), **kw)


# ------------------------------------------------------------------ settings
def test_older_documents_get_the_new_settings(store):
    d = store.create_document("Old", {})
    raw = json.loads(store._one("SELECT settings FROM documents WHERE id = ?", (d["id"],))["settings"])
    for k in ("cover", "style", "header", "footer", "back"):
        raw.pop(k, None)
    store._db.execute("UPDATE documents SET settings = ? WHERE id = ?", (json.dumps(raw), d["id"]))
    s = store.get_document(d["id"])["settings"]
    assert s["cover"]["enabled"] is True and s["style"] == {} and s["footer"]["left"] == ""


def test_style_overrides_validate_merge_and_reset(store):
    d = store.create_document("S", {})
    s = store.update_document(d["id"], settings={"style": {"color_accent": "#002f58", "font_size": 14}})["settings"]
    assert s["style"] == {"color_accent": "#002f58", "font_size": 14.0}
    s = store.update_document(d["id"], settings={"style": {"color_accent": None, "heading_case": "uppercase"}})["settings"]
    assert s["style"] == {"font_size": 14.0, "heading_case": "uppercase"}
    assert store.update_document(d["id"], settings={"style": None})["settings"]["style"] == {}
    for bad in ({"color_accent": "red; x"}, {"font_body": "a{b}"}, {"font_size": 99}, {"nope": 1}):
        with pytest.raises(Invalid):
            store.update_document(d["id"], settings={"style": bad})


def test_cover_and_header_footer_validation(store):
    d = store.create_document("C", {})
    s = store.update_document(d["id"], settings={"cover": {"enabled": False, "background": "#002f58", "align": "top"},
                                                 "footer": {"right": "{page}", "rule": True}})["settings"]
    assert s["cover"] == {"enabled": False, "align": "top", "background": "#002f58", "image": None, "text": None}
    assert s["footer"]["right"] == "{page}" and s["footer"]["rule"] is True and s["header"]["left"] == ""
    for bad in ({"cover": {"align": "left"}}, {"cover": {"background": "url(x)"}}, {"cover": {"image": "javascript:x"}},
                {"back": {"text": "loud"}}, {"header": {"left": "x" * 400}}):
        with pytest.raises(Invalid):
            store.update_document(d["id"], settings=bad)


def test_effective_theme_layers_the_documents_overrides():
    t = themes.effective_theme({"theme": "modern", "style": {"color_accent": "#123456"}})
    assert t["color_accent"] == "#123456" and t["font_body"].startswith("Inter")
    assert themes.effective_theme({"theme": "gone"})["label"] == "Classic"


# ------------------------------------------------------------------ cover role
def test_blocks_added_after_a_cover_block_stay_on_the_cover(store):
    d = _report(store)
    title, sub = store.list_blocks(d["id"])[:2]
    assert store.insert_block(d["id"], "text", {"markdown": "x"}, after=sub["id"])["layout"]["role"] == "cover"
    assert store.insert_block(d["id"], "text", {"markdown": "x"}, after=sub["id"],
                              layout={"role": "flow"})["layout"]["role"] == "flow"
    assert store.insert_block(d["id"], "text", {"markdown": "y"})["layout"]["role"] == "flow"
    sg = store.propose(d["id"], "insert", {"type": "image", "content": {}, "after": title["id"], "layout": {"span": "full"}})
    assert sg["payload"]["layout"] == {"span": "full", "role": "cover"}


def test_propose_insert_takes_a_role(store, cfg_for):
    d = store.create_document("R", {})
    out = propose_insert.invoke({"block_type": "heading", "content": "Back matter", "role": "back"}, config=cfg_for(d["id"]))
    assert "sug_" in out
    assert store.list_suggestions(d["id"])[0]["payload"]["layout"]["role"] == "back"


# ------------------------------------------------------------------ agent tools
def test_design_tools(store, cfg_for):
    d = _report(store)
    c = cfg_for(d["id"])
    out = set_typography.invoke({"font_body": "Lato", "color_heading": "#002f58", "font_size_px": 14}, config=c)
    assert "'Lato'" in out
    style = store.get_document(d["id"])["settings"]["style"]
    assert style["font_body"].startswith("'Lato', ") and style["color_heading"] == "#002f58"
    assert "ERROR" in set_typography.invoke({"color_text": "blue;"}, config=c)
    assert set_typography.invoke({"reset": True}, config=c) == "Typography reset to the theme."
    assert store.get_document(d["id"])["settings"]["style"] == {}

    set_cover.invoke({"enabled": False}, config=c)
    assert store.get_document(d["id"])["settings"]["cover"]["enabled"] is False
    assert "OFF" in get_document.invoke({}, config=c)
    set_cover.invoke({"enabled": True, "background": "#002f58", "text_color": "light", "align": "bottom"}, config=c)
    assert "background #002f58" in get_document.invoke({}, config=c)
    assert "no block has role=back" in set_cover.invoke({"page": "back", "background": "#000000"}, config=c)

    out = set_header_footer.invoke({"footer_left": "{title}", "footer_right": "{page}", "first_page": False}, config=c)
    assert "Footer: {title} | — | {page} (not on the first body page)" in out
    assert "Header and footer removed." == set_header_footer.invoke({"clear": True}, config=c)


def test_match_upload_style_tool_reads_a_word_file(store, cfg_for):
    d = store.create_document("M", {})
    u = store.add_upload(d["id"], "brand.docx", _docx(), purpose="reference",
                         mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    out = match_upload_style.invoke({"upload_id": u["id"]}, config=cfg_for(d["id"]))
    assert "Calibri → Carlito" in out and "#0b5d3b" in out, out
    s = store.get_document(d["id"])["settings"]
    assert s["page_size"] == "a4" and s["columns"] == 2 and s["style"]["color_heading"] == "#0b5d3b"


# ------------------------------------------------------------------ fonts
@pytest.mark.parametrize("raw,family,exact", [
    ("ABCDEF+TimesNewRomanPSMT", "Tinos", False),
    ("Calibri-Bold", "Carlito", False),
    ("ABCDEF+Raleway-SemiBold", "Raleway", True),
    ("Montserrat-ExtraBold", "Montserrat", True),
    ("PlayfairDisplay-Bold", "Playfair Display", True),
    ("HelveticaNeue-Medium", "Inter", False),
    ("SourceSansPro-Regular", "Source Sans 3", False),
    ("MysteryFaceSerif", "Source Serif 4", False),
])
def test_font_mapping(raw, family, exact):
    m = fonts.map_font(raw)
    assert (m["family"], m["exact"]) == (family, exact)


def test_bundled_fonts_match_the_frontend():
    from pathlib import Path

    src = (Path(__file__).resolve().parents[1] / "frontend" / "src" / "fonts.ts").read_text(encoding="utf-8")
    for fam in fonts.BUNDLED:
        assert f'"{fam}"' in src, fam


# ------------------------------------------------------------------ style reading
def _docx() -> bytes:
    import docx
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Mm(25), Mm(20), Mm(22), Mm(22)
    cols = sec._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = sec._sectPr.makeelement(qn("w:cols"), {})
        sec._sectPr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "567")
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    for name, size in (("Heading 1", 18), ("Heading 2", 14)):
        st = d.styles[name]
        st.font.name, st.font.size = "Cambria", Pt(size)
        st.font.color.rgb = RGBColor(0x0B, 0x5D, 0x3B)
    d.add_heading("Catchment report", 1)
    d.add_paragraph("Body text.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_word_style_profile():
    prof = stylematch.analyze("brand.docx", _docx(), "application/octet-stream", [])
    assert prof["page"]["size"] == "a4" and prof["columns"] == 2 and prof["column_gap"] == 10.0
    assert prof["margins"] == {"top": 25.0, "right": 22.0, "bottom": 20.0, "left": 22.0}
    assert prof["fonts"]["body"]["family"] == "Carlito" and prof["fonts"]["heading"]["family"] == "Caladea"
    assert prof["colors"]["text"] == "#333333" and prof["colors"]["heading"] == "#0b5d3b"
    patch, changes = stylematch.to_settings(prof)
    assert patch["style"]["font_size"] == 14.0 and patch["style"]["heading_scale"] == pytest.approx(1.155, abs=0.01)
    assert any(c.startswith("Fonts: Calibri → Carlito") for c in changes)


def test_page_image_profile(tmp_path):
    from PIL import Image, ImageDraw

    w, h = 794, 1123  # A4 at 96 dpi
    img = Image.new("RGB", (w, h), "#ffffff")
    dr = ImageDraw.Draw(img)
    for col in (76, 420):  # two columns of "text" bars, 76px margins, a 46px gutter
        for y in range(95, 1020, 18):
            dr.rectangle([col, y, col + 298, y + 9], fill="#222222")
    dr.rectangle([76, 600, 374, 700], fill="#e6f5fc")
    path = tmp_path / "page.png"
    img.save(path)
    prof = stylematch.analyze("page.png", b"", "image/png", [str(path)])
    assert prof["page"]["size"] == "a4" and prof["page"]["guessed"]
    assert prof["columns"] == 2
    assert prof["colors"]["tint"] == "#e6f5fc" and prof["colors"]["text"] == "#222222"
    assert abs(prof["margins"]["left"] - 20.1) < 1.5


# ------------------------------------------------------------------ API
def test_house_defaults_saved_themes_and_config(tmp_path):
    app = _app(tmp_path, doc_defaults={"cover": {"enabled": False}, "footer": {"left": "{page}", "right": "SWAN Forum"}})
    with TestClient(app) as c:
        cfg = c.get("/api/config").json()
        assert cfg["cover_default"] is False and cfg["web_fonts"] is True
        assert any(f["family"] == "Raleway" and f["bundled"] for f in cfg["fonts"])
        assert next(t for t in cfg["templates"] if t["id"] == "report")["has_cover"] is True
        doc = c.post("/api/docs", json={"template": "report", "title": "Pilot"}).json()
        assert doc["settings"]["cover"]["enabled"] is False and doc["settings"]["footer"]["right"] == "SWAN Forum"
        on = c.post("/api/docs", json={"template": "report", "settings": {"cover": {"enabled": True}}}).json()
        assert on["settings"]["cover"]["enabled"] is True

        c.patch(f"/api/docs/{doc['id']}", json={"settings": {"style": {"color_accent": "#009fe3"},
                                                                  "cover": {"background": "#002f58"}}})
        r = c.post("/api/themes", json={"label": "Example Forum", "doc_id": doc["id"]}).json()
        assert r["name"] == "example-forum"
        saved = r["themes"]["example-forum"]
        assert saved["saved"] and saved["color_accent"] == "#009fe3" and saved["cover_background"] == "#002f58"
        assert c.post("/api/themes", json={"label": "Example Forum", "values": {}}).json()["name"] == "example-forum-2"
        assert c.post("/api/themes", json={"label": "", "values": {}}).status_code == 400
        assert c.patch(f"/api/docs/{on['id']}", json={"settings": {"theme": "example-forum"}}).status_code == 200
        assert c.delete("/api/themes/example-forum-2").status_code == 200
        assert c.delete("/api/themes/classic").status_code == 404
    themes.set_saved_provider(None)


def test_style_upload_and_apply(tmp_path):
    app = _app(tmp_path)
    with TestClient(app) as c:
        doc = c.post("/api/docs", json={"template": "blank"}).json()
        u = c.post(f"/api/docs/{doc['id']}/uploads", data={"purpose": "style"},
                   files={"file": ("brand.docx", _docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}).json()
        assert u["style"]["source"] == "docx"
        r = c.post(f"/api/uploads/{u['id']}/apply-style").json()
        assert r["doc"]["settings"]["columns"] == 2 and r["doc"]["settings"]["style"]["color_heading"] == "#0b5d3b"
        assert any("Carlito" in ch for ch in r["changes"])
        ref = c.post(f"/api/docs/{doc['id']}/uploads", data={"purpose": "reference"},
                     files={"file": ("notes.txt", b"plain", "text/plain")}).json()
        assert c.post(f"/api/uploads/{ref['id']}/apply-style").status_code == 400


def test_pdf_uploads_get_rendered_pages(tmp_path):
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=595, height=842)
    w.add_blank_page(width=595, height=842)
    buf = io.BytesIO()
    w.write(buf)
    app = _app(tmp_path)
    with TestClient(app) as c:
        doc = c.post("/api/docs", json={"template": "blank"}).json()
        u = c.post(f"/api/docs/{doc['id']}/uploads", data={"purpose": "reference"},
                   files={"file": ("blank.pdf", buf.getvalue(), "application/pdf")}).json()
        assert len(u["pages"]) == 2 and all(p.endswith(".jpg") for p in u["pages"])
        c.delete(f"/api/uploads/{u['id']}")
        from pathlib import Path

        assert not any(Path(p).exists() for p in u["pages"])


# ------------------------------------------------------------------ Word export
def test_docx_export_carries_the_design(store):
    import docx

    from langstage_codraft import exporters

    d = store.create_document("Pilot", {"cover": {"enabled": False},
                                        "style": {"font_body": "'Lato', sans-serif", "color_heading": "#002f58"},
                                        "footer": {"left": "{page}", "center": "{title}", "rule": True}},
                              tpl.BUILTIN_TEMPLATES["report"]["blocks"])
    x = docx.Document(io.BytesIO(exporters.to_docx(store, d["id"], {}, [])))
    assert len(x.sections) == 1  # no cover page: no section break
    assert x.styles["Normal"].font.name == "Lato" and str(x.styles["Heading 2"].font.color.rgb) == "002F58"
    foot = x.sections[0].footer.paragraphs[0]
    assert "Pilot" in foot.text and "PAGE" in foot._p.xml and "pBdr" in foot._p.xml
    store.update_document(d["id"], settings={"cover": {"enabled": True}})
    x = docx.Document(io.BytesIO(exporters.to_docx(store, d["id"], {}, [])))
    assert len(x.sections) == 2 and not x.sections[0].footer.paragraphs[0].text


# ------------------------------------------------------------------ the agent's name
def _named_graph(name=None):
    from langchain.agents import create_agent
    from langchain_core.language_models.fake_chat_models import FakeListChatModel

    return create_agent(FakeListChatModel(responses=["ok"]), name=name)


@pytest.mark.parametrize("graph_name,configured,shown", [
    ("Juniper", None, "Juniper"),          # the graph's own name
    (None, None, "Co-author"),             # LangGraph's default name isn't shown
    ("Juniper", "Acme Assistant", "Acme Assistant"),  # a deployer's explicit choice wins
])
def test_the_editor_calls_the_agent_by_its_graph_name(tmp_path, graph_name, configured, shown):
    over = {"data_dir": tmp_path / "data", "port": 8796, **({"agent_name": configured} if configured else {})}
    cfg = CodraftConfig.resolve(overrides=over, use_toml=False)
    info = {"kind": "spec", "label": "my_agent.py:graph", "missing_tools": []}
    app = create_app(cfg, _named_graph(graph_name), info, store=Store(tmp_path / "data" / "codraft.db"))
    with TestClient(app) as c:
        conf = c.get("/api/config").json()
    assert conf["ui"]["agent_name"] == shown
    if shown != "Co-author":  # the status chip shows the name, keeping the spec alongside
        assert conf["agent"]["label"] == shown and conf["agent"]["spec"] == "my_agent.py:graph"
    else:
        assert conf["agent"]["label"] == "my_agent.py:graph"
